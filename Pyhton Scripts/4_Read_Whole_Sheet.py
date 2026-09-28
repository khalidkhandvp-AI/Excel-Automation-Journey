from openpyxl import load_workbook
data1=load_workbook("Excel files/Student File.xlsx")
print(data1.sheetnames)
Sheet1=data1.active
for row in Sheet1.iter_rows():
    for data in row:
        print(data.value)
        