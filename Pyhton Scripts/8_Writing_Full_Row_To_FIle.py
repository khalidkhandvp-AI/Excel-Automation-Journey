from openpyxl import load_workbook
wb=load_workbook("Excel Files/Student File.xlsx")
print(wb.sheetnames)
Sheet=wb["Sheet1"]
data=["khan1","khan2","khan3","khan4","khan5"]
for column_value,entry in enumerate(data,start=1):
    Sheet.cell(row=21,column=column_value,value=entry)
    wb.save("Excel Files/Student File.xlsx")