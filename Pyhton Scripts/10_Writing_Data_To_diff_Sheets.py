from openpyxl import load_workbook
wb=load_workbook("Excel Files/Student File.xlsx")
shet=wb["Sheet1"]
sheet=wb["Sheet2"]
sheeet=wb["Sheet3"]
shet.cell(row=5,column=9,value="Khalid")
sheet.cell(row=5,column=9,value="KHan")
sheeet.cell(row=5,column=9,value="Buner")
wb.save("Excel Files/Student File.xlsx")