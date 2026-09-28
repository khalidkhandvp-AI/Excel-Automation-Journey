from openpyxl import load_workbook
data=load_workbook("Excel Files/Student File.xlsx")
print(data.sheetnames)
Sheet1=data.active
# A SINGLE ROW
for row in Sheet1.iter_rows(min_row=4,max_row=4,min_col=1,max_col=6):
    for cell in row:
        print(cell.value)
#MORE ROWS
for row in Sheet1.iter_rows(min_row=3,max_row=6,min_col=1,max_col=6):
    for cell in row:
        print(cell.value)