from openpyxl import load_workbook
workbook=load_workbook('Excel Files/Student File.xlsx')
print(workbook.sheetnames)
Sheet1=workbook['Sheet1']
Sheet2=workbook['Sheet2']
Sheet3=workbook['Sheet3']
print(Sheet1.cell(row=4,column=1).value)
print(Sheet2.cell(row=4,column=1).value)
print(Sheet3.cell(row=4,column=1).value)