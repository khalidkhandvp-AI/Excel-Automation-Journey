from openpyxl import load_workbook
wb=load_workbook("Excel Files/Student File.xlsx")
ws=wb["Sheet2"]
data=["Player","Player","Player","Player","Player","Player","Player","Player","Player","Player","Player",]
ws.insert_cols(1)
for row_id,entry in enumerate(data,start=1):
    ws.cell(row=row_id,column=1,value=entry)
wb.save("Excel Files/Student File.xlsx")