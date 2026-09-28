from openpyxl import load_workbook
from openpyxl.styles import Font
wb=load_workbook("Excel Files/Student File.xlsx")
ws=wb["Sheet1"]
ws["B21"].font= Font(
    name="Arial",
    size=20,
    bold=True
)
wb.save("Excel Files/Student File.xlsx")