from openpyxl import load_workbook
wb=load_workbook("Excel Files/Student File.xlsx")
sheet=wb["Sheet1"]
data=[1,2,3,4,5,6,7,8,9,10,11,12,13,14]
for row_id,entry in enumerate(data,start=1):
    sheet.cell(row=row_id,column=10,value=entry)
    wb.save("Excel Files/Student File.xlsx")