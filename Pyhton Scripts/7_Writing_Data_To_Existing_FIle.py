from openpyxl import load_workbook
workbook=load_workbook("Excel Files/Student file.xlsx")
Sheet1=workbook.active
Sheet1.cell(row=20,column=10,value="filled")
workbook.save("Excel Files/Student file.xlsx")
workbook.close