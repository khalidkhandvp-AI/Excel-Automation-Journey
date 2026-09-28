from openpyxl import load_workbook
wb=load_workbook("Excel Files/Student File.xlsx")
ws=wb["Sheet2"]
#INSERTING AN EMPTY ROW 
ws.insert_rows(3)
wb.save("Excel Files/Student File.xlsx")
#INSERTING DATA INTO THE ROW
data=["Khalid","Khan"]
for colm,valu in enumerate(data,start=1):
    ws.cell(row=3,column=colm,value=valu)
wb.save("Excel Files/Student File.xlsx")