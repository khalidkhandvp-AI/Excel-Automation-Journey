from openpyxl import load_workbook
data=load_workbook("Excel Files/Student File.xlsx")
print(data.sheetnames)
Sheet1=data.active
#SINGLE COLUMN
for col in Sheet1.iter_cols(min_row=1,max_row=10,min_col=1,max_col=1):
    for cell in col:
        print(cell.value)
#MORE COLUMNS
for col in Sheet1.iter_cols(min_row=1,max_row=10,min_col=1,max_col=3):
    for cell in col:
        print(cell.value)