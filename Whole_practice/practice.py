#FILE CREATION
from openpyxl import Workbook
wb=Workbook()
ws=wb.active
ws.title="Employees"
wb.save("employees_data.xlsx")
print("file created")


#HEADERS WRITING
HEADERS=['ID','Name','Department','Salary','City']
for col,header in enumerate(HEADERS,start=1):
    ws.cell(row=1,column=col,value=header)
    wb.save("employees_data.xlsx")


#WRITING DATA TO FILE
data= [
[1,"Ali","IT",50000,"Peshawar"],
[2,"Ahmad","HR",45000,"Mardan"],
[3,"Usman","Finance",55000,"Swat"],
[4,"Bilal","IT",60000,"Charsadda"],
[5,"Hamza","Marketing",47000,"Nowshera"],
[6,"Saad","HR",49000,"Peshawar"],
[7,"Zain","IT",62000,"Mardan"],
[8,"Ayan","Finance",53000,"Swat"],
[9,"Talha","Marketing",51000,"Charsadda"],
[10,"Hassan","IT",65000,"Peshawar"],
[11,"Farhan","HR",46000,"Nowshera"],
[12,"Rizwan","Finance",58000,"Mardan"],
[13,"Junaid","Marketing",50000,"Swat"],
[14,"Imran","IT",70000,"Peshawar"],
[15,"Khalid","HR",48000,"Charsadda"],
[16,"Asad","Finance",56000,"Mardan"],
[17,"Shayan","IT",63000,"Swat"],
[18,"Arham","Marketing",52000,"Nowshera"],
[19,"Daniyal","HR",47000,"Peshawar"],
[20,"Noman","IT",68000,"Mardan"],
[21,"Ibrahim","Finance",59000,"Swat"],
[22,"Haris","Marketing",54000,"Charsadda"],
[23,"Waqas","HR",50000,"Nowshera"],
[24,"Shahzaib","IT",72000,"Peshawar"],
[25,"Rehan","Finance",61000,"Mardan"],
[26,"Fahad","Marketing",53000,"Swat"],
[27,"Adnan","HR",49000,"Charsadda"],
[28,"Yasir","IT",66000,"Nowshera"],
[29,"Jawad","Finance",57000,"Peshawar"],
[30,"Anas","Marketing",55000,"Mardan"],
[31,"Taimoor","HR",51000,"Swat"],
[32,"Adeel","IT",74000,"Charsadda"],
[33,"Sameer","Finance",62000,"Nowshera"],
[34,"Qasim","Marketing",56000,"Peshawar"],
[35,"Mudassir","HR",52000,"Mardan"],
[36,"Sami","IT",69000,"Swat"],
[37,"Naveed","Finance",60000,"Charsadda"],
[38,"Omer","Marketing",54000,"Nowshera"],
[39,"Kamran","HR",50000,"Peshawar"],
[40,"Huzaifa","IT",76000,"Mardan"],
[41,"Basit","Finance",63000,"Swat"],
[42,"Aqib","Marketing",57000,"Charsadda"],
[43,"Ahsan","HR",53000,"Nowshera"],
[44,"Taha","IT",71000,"Peshawar"],
[45,"Salman","Finance",64000,"Mardan"],
[46,"Rafay","Marketing",58000,"Swat"],
[47,"Umair","HR",52000,"Charsadda"],
[48,"Sufyan","IT",73000,"Nowshera"],
[49,"Shakir","Finance",65000,"Peshawar"],
[50,"Hammad","Marketing",59000,"Mardan"]
]
for row in data:
    ws.append(row)
wb.save("employees_data.xlsx")
print("insert done")


#Reading data from cell
from openpyxl import load_workbook
wb=load_workbook("employees_data.xlsx")
ws=wb.active
Value_c10=ws.cell(row=10,column=3).value
print(Value_c10)


#Reading complete row from sheet
for row in ws.iter_rows(min_row=3,max_row=3,min_col=1,max_col=6):
    for cell in row:
        print(cell.value)


#Reading column from sheet
for col in ws.iter_cols(min_row=1,max_row=50,min_col=2,max_col=2):
    for cell in col:
        print(cell.value)


#Reading complete excel sheet
for row in ws.iter_rows():
    for data in row:
        print(data.value)

#Writing value to specific cell two ways
ws.cell(row=53,column=1,value="KHAN KHAN")
wb.save("employees_data.xlsx")
ws["A54"]="KHALID"
wb.save("employees_data.xlsx")

#Writing entire row to a sheet
data2=[51,'fahim','AI','850000','Buner']
for column_id,entry in enumerate(data2,start=1):
    ws.cell(row=52,column=column_id,value=entry)
wb.save("employees_data.xlsx")


#Writing entire column to a sheet
data3=["COntact",
    "03014827156",
    "03051938472",
    "03086742195",
    "03028574136",
    "03072918453",
    "03105739184",
    "03128462715",
    "03152946837",
    "03187154926",
    "03116385274",
    "03218471536",
    "03242958174",
    "03276149285",
    "03293817462",
    "03227591843",
    "03315847291",
    "03339174628",
    "03352618497",
    "03387451926",
    "03304829617",
    "03417582914",
    "03431948572",
    "03458264719",
    "03485719246",
    "03402936847",
    "03037462918",
    "03061857294",
    "03096248175",
    "03139572184",
    "03164827591",
    "03192746185",
    "03208619472",
    "03234157829",
    "03267931845",
    "03285462917",
    "03321847596",
    "03346729185",
    "03364295718",
    "03378152946",
    "03392617485",
    "03429751846",
    "03443185729",
    "03467524918",
    "03471946285",
    "03498261574",
    "03045718462",
    "03142685197",
    "03257149286",
    "03304918572",
    "03456382719"
]
for row_id,entry in enumerate(data3,start=1):
    ws.cell(row=row_id,column=6,value=entry)
wb.save("employees_data.xlsx")


#writing data to different sheet of same excel file
ws2=wb.create_sheet("Products")
names=["Product ID","Name","Category","Price","Quantity Sold","Revenue","Rating"]
for col,heade in enumerate (names,start=1):
    ws2.cell(row=1,column=col,value=heade)
wb.save("employees_data.xlsx")
data4= [
    ["ProductID", "ProductName", "Category", "Price", "QuantitySold", "Revenue", "Rating"],
    [1, "Laptop", "Electronics", 85000, 12, 1020000, 4.7],
    [2, "Mouse", "Electronics", 1200, 45, 54000, 4.3],
    [3, "Keyboard", "Electronics", 2500, 30, 75000, 4.4],
    [4, "Headphones", "Electronics", 3500, 25, 87500, 4.5],
    [5, "Monitor", "Electronics", 28000, 10, 280000, 4.6],
    [6, "Printer", "Electronics", 22000, 8, 176000, 4.2],
    [7, "USB Drive", "Accessories", 1500, 40, 60000, 4.1],
    [8, "Power Bank", "Accessories", 3000, 20, 60000, 4.4],
    [9, "Smart Watch", "Wearables", 12000, 15, 180000, 4.6],
    [10, "Tablet", "Electronics", 45000, 7, 315000, 4.5],
    [11, "Speaker", "Electronics", 5500, 18, 99000, 4.3],
    [12, "Router", "Networking", 6500, 14, 91000, 4.2],
    [13, "Webcam", "Accessories", 4000, 11, 44000, 4.1],
    [14, "Microphone", "Accessories", 7000, 9, 63000, 4.5],
    [15, "SSD 512GB", "Storage", 9500, 16, 152000, 4.8],
    [16, "HDD 1TB", "Storage", 7500, 13, 97500, 4.4],
    [17, "Graphics Card", "Components", 95000, 4, 380000, 4.9],
    [18, "RAM 16GB", "Components", 14000, 12, 168000, 4.7],
    [19, "CPU", "Components", 55000, 6, 330000, 4.8],
    [20, "Motherboard", "Components", 25000, 8, 200000, 4.6]
]
for ro in data4:
    ws2.append(ro)
wb.save("employees_data.xlsx")


#Replacing cell value 
ws.cell(row=2,column=5,value="BUNER")
wb.save("employees_data.xlsx")


#Formating cell values
from openpyxl.styles import Font
for cell in ws[1]:
    cell.font=Font(bold=True)
wb.save("employees_data.xlsx")


#Coustomizing font type and size
from openpyxl.styles import Font
ws=wb.active
cell_a53=ws.cell(row=53,column=1)
font_a1=Font(name="Ahroni",size=15)
cell_a53.font=font_a1
wb.save("employees_data.xlsx")

#changing cell backround colour
from openpyxl.styles import PatternFill
ws['A53'].fill=PatternFill(
    fill_type='solid',
    fgColor="FFFF00"
)
wb.save("employees_data.xlsx")


#changing cell text colour
#ws['A54'].font=Font(
#   fill_type='solid',
#  fgColor="FF0000"
#)
#wb.save("employees_data.xlsx")


#How to insert a blank row 
ws.insert_rows(5)
#Now putting values in the cells of that row
ws['A5']="abdul"
ws['B5']='AI'
ws['C5']='850000'
ws['D5']='Buner'
ws['E5']='03327843759'
wb.save("employees_data.xlsx")
#Now deleting the row
ws.delete_rows(5)
wb.save("employees_data.xlsx")


#How to insert a blank column
ws.insert_cols(3)
wb.save("employees_data.xlsx")
#putting values to them
ws['c1']='amn'
ws['c2']='kas'
ws['c3']='MKK'
ws['c4']='jkl'
wb.save("employees_data.xlsx")
#now delete that column
ws.delete_cols(3)
wb.save("employees_data.xlsx")

#Applying formulas in different columns
formual='=SUM(D2:D52)'
ws['D53']=formual
wb.save("employees_data.xlsx")